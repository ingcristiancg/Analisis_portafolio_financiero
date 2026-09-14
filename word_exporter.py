"""
Generador de Informes Ejecutivos en Word (.docx)
Exporta el reporte cuantitativo completo:
- Portada y resumen ejecutivo.
- Tablas estadísticas formateadas (parámetros mensuales estrictos).
- Gráficos de alta resolución generados con Matplotlib.
- Dictamen y recomendaciones del Agente Cuantitativo Financiero Autónomo.
"""

from typing import Dict, Any, Optional
import io
import os
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import seaborn as sns
import numpy as np
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn


def set_cell_background(cell, fill_hex: str):
    """Aplica color de fondo a una celda de Word."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Ajusta márgenes internos de una celda."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)


def create_charts_images(
    stats_data: Dict[str, Any],
    opt_data: Dict[str, Any],
    mc_data: Dict[str, Any],
    bench_data: Optional[Dict[str, Any]],
) -> Dict[str, io.BytesIO]:
    """
    Genera gráficos en alta resolución para incrustar en el documento Word.
    """
    charts = {}
    sns.set_theme(style="whitegrid", palette="deep")

    # 1. Gráfico de Monte Carlo y Frontera Eficiente
    fig, ax = plt.subplots(figsize=(8.5, 5.0), dpi=300)
    sim_df = mc_data["simulation_df"]
    
    scatter = ax.scatter(
        sim_df["Volatilidad Mensual"] * 100,
        sim_df["Rendimiento Mensual"] * 100,
        c=sim_df["Ratio Sharpe Mensual"],
        cmap="viridis",
        alpha=0.6,
        s=18,
        edgecolors="none",
    )
    cbar = plt.colorbar(scatter, ax=ax)
    cbar.set_label("Ratio de Sharpe Mensual (Rf=0)", fontsize=10, fontweight="bold")

    # Cartera Óptima
    opt_vol = opt_data["volatility"] * 100
    opt_ret = opt_data["expected_return"] * 100
    ax.scatter(
        [opt_vol],
        [opt_ret],
        color="#E63946",
        s=180,
        marker="*",
        label=f"Cartera Óptima (Ret: {opt_ret:.2f}%, Vol: {opt_vol:.2f}%)",
        zorder=5,
        edgecolors="black",
    )

    # Línea de restricción de riesgo (sigma <= 7%)
    ax.axvline(
        x=opt_data["target_max_std"] * 100,
        color="#D62828",
        linestyle="--",
        linewidth=1.5,
        label=f"Tope Máx. Riesgo Mensual (≤ {opt_data['target_max_std']*100:.1f}%)",
    )

    # Benchmark IPC si existe
    if bench_data:
        ipc_vol = bench_data["ipc_std_monthly"] * 100
        ipc_ret = bench_data["ipc_mean_monthly"] * 100
        ax.scatter(
            [ipc_vol],
            [ipc_ret],
            color="#2A9D8F",
            s=120,
            marker="D",
            label=f"Benchmark IPC (Ret: {ipc_ret:.2f}%, Vol: {ipc_vol:.2f}%)",
            zorder=4,
            edgecolors="black",
        )

    ax.set_title("Simulación de Monte Carlo (1,000 Carteras) y Cartera Óptima de Markowitz", fontsize=12, fontweight="bold", pad=12)
    ax.set_xlabel("Volatilidad Mensual σ_p (%) [PROHIBIDO ANUALIZAR]", fontsize=10, fontweight="bold")
    ax.set_ylabel("Rendimiento Esperado Mensual μ_p (%)", fontsize=10, fontweight="bold")
    ax.legend(loc="upper left", frameon=True, fontsize=9)
    plt.tight_layout()

    buf_mc = io.BytesIO()
    plt.savefig(buf_mc, format="png", bbox_inches="tight")
    buf_mc.seek(0)
    charts["monte_carlo"] = buf_mc
    plt.close(fig)

    # 2. Gráfico de Ponderaciones Óptimas
    fig2, ax2 = plt.subplots(figsize=(7.5, 4.2), dpi=300)
    w_df = opt_data["weights_df"].copy()
    w_active = w_df[w_df["Ponderación Óptima (w)"] > 0.001].copy()
    
    colors = sns.color_palette("Blues_r", n_colors=len(w_active))
    bars = ax2.barh(w_active["Activo"], w_active["Porcentaje (%)"], color=colors, edgecolor="#1D3557", height=0.55)
    
    for bar in bars:
        width = bar.get_width()
        ax2.text(
            width + 0.6,
            bar.get_y() + bar.get_height() / 2,
            f"{width:.2f}%",
            ha="left",
            va="center",
            fontsize=10,
            fontweight="bold",
            color="#1D3557",
        )

    ax2.set_xlim(0, max(w_active["Porcentaje (%)"]) * 1.18)
    ax2.invert_yaxis()
    ax2.set_title("Asignación Óptima de Capital (Teoría de Markowitz - Posiciones Largas)", fontsize=12, fontweight="bold", pad=12)
    ax2.set_xlabel("Ponderación en el Portafolio (%)", fontsize=10, fontweight="bold")
    plt.tight_layout()

    buf_weights = io.BytesIO()
    plt.savefig(buf_weights, format="png", bbox_inches="tight")
    buf_weights.seek(0)
    charts["weights"] = buf_weights
    plt.close(fig2)

    # 3. Matriz de Correlación
    fig3, ax3 = plt.subplots(figsize=(7.5, 6.0), dpi=300)
    corr = stats_data["corr_matrix"]
    sns.heatmap(
        corr,
        annot=True,
        fmt=".2f",
        cmap="coolwarm",
        vmin=-1,
        vmax=1,
        ax=ax3,
        cbar_kws={"label": "Coeficiente de Correlación de Pearson"},
        annot_kws={"size": 8},
        linewidths=0.5,
    )
    ax3.set_title("Matriz de Correlación de Rendimientos Mensuales", fontsize=12, fontweight="bold", pad=12)
    plt.xticks(rotation=45, ha="right", fontsize=9)
    plt.yticks(rotation=0, fontsize=9)
    plt.tight_layout()

    buf_corr = io.BytesIO()
    plt.savefig(buf_corr, format="png", bbox_inches="tight")
    buf_corr.seek(0)
    charts["correlation"] = buf_corr
    plt.close(fig3)

    return charts


def generate_word_report(
    prep_data: Dict[str, Any],
    stats_data: Dict[str, Any],
    opt_data: Dict[str, Any],
    mc_data: Dict[str, Any],
    bench_data: Optional[Dict[str, Any]],
    agent_report: Dict[str, Any],
    output_path: Optional[str] = None,
) -> io.BytesIO:
    """
    Genera el documento Word (.docx) formal y estructurado con todos los hallazgos.
    """
    doc = Document()

    # Configuración de página
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Paleta de colores corporativos
    HEX_NAVY = "1B365D"
    HEX_GOLD = "D4AF37"
    HEX_BG_LIGHT = "F4F6F9"
    # ----------------------------------------------------
    # PORTADA ACADÉMICA OFICIAL (portada.docx)
    # ----------------------------------------------------
    logo_path = os.path.join("assets", "logo_universidad.png")
    if not os.path.exists(logo_path):
        logo_path = os.path.join("assets", "logo_universidad.jpg")

    if os.path.exists(logo_path):
        p_logo = doc.add_paragraph()
        p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_logo.paragraph_format.space_before = Pt(30)
        p_logo.paragraph_format.space_after = Pt(25)
        run_logo = p_logo.add_run()
        run_logo.add_picture(logo_path, width=Inches(2.4))

    p_tr = doc.add_paragraph()
    p_tr.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tr.paragraph_format.space_before = Pt(15)
    p_tr.paragraph_format.space_after = Pt(6)
    r_tr = p_tr.add_run("TRABAJO")
    r_tr.font.name = "Times New Roman"
    r_tr.font.size = Pt(14)
    r_tr.font.bold = True

    p_tit = doc.add_paragraph()
    p_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_tit.paragraph_format.space_after = Pt(35)
    r_tit = p_tit.add_run("Portafolio diversificado usando rendimientos reales")
    r_tit.font.name = "Times New Roman"
    r_tit.font.size = Pt(14)
    r_tit.font.bold = True

    p_pres = doc.add_paragraph()
    p_pres.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_pres.paragraph_format.space_after = Pt(4)
    r_pres = p_pres.add_run("Presentan:")
    r_pres.font.name = "Times New Roman"
    r_pres.font.size = Pt(12)

    p_aut = doc.add_paragraph()
    p_aut.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut.paragraph_format.space_after = Pt(30)
    r_aut = p_aut.add_run("Cristian Andres Cordoba Gonzalez (Colombia)")
    r_aut.font.name = "Times New Roman"
    r_aut.font.size = Pt(12)
    r_aut.font.bold = True

    p_prof_lbl = doc.add_paragraph()
    p_prof_lbl.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prof_lbl.paragraph_format.space_after = Pt(4)
    r_prof_lbl = p_prof_lbl.add_run("Profesor:")
    r_prof_lbl.font.name = "Times New Roman"
    r_prof_lbl.font.size = Pt(12)

    p_prof_val = doc.add_paragraph()
    p_prof_val.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_prof_val.paragraph_format.space_after = Pt(35)
    r_prof_val = p_prof_val.add_run("Daniel Vázquez Cotera")
    r_prof_val.font.name = "Times New Roman"
    r_prof_val.font.size = Pt(12)
    r_prof_val.font.bold = True

    p_univ = doc.add_paragraph()
    p_univ.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_univ.paragraph_format.space_after = Pt(6)
    r_univ = p_univ.add_run("Broward International University")
    r_univ.font.name = "Times New Roman"
    r_univ.font.size = Pt(13)
    r_univ.font.bold = True

    p_mba = doc.add_paragraph()
    p_mba.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_mba.paragraph_format.space_after = Pt(35)
    r_mba = p_mba.add_run("Financial Management with Artificial Intelligence MBA")
    r_mba.font.name = "Times New Roman"
    r_mba.font.size = Pt(12)

    p_date = doc.add_paragraph()
    p_date.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_date = p_date.add_run("13 de septiembre de 2026")
    r_date.font.name = "Times New Roman"
    r_date.font.size = Pt(11)

    doc.add_page_break()

    # 1. Encabezado principal
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    run_title = title_p.add_run("INFORME CUANTITATIVO DE GESTIÓN Y OPTIMIZACIÓN DE CARTERAS")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(20)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    run_sub = sub_p.add_run("Aplicación Estricta de la Teoría de Markowitz y Simulación de Monte Carlo (Datos Mensuales)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

    # 2. Caja de Metadatos del Informe
    meta_table = doc.add_table(rows=2, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_table.autofit = False

    col_widths = [Inches(3.3), Inches(3.3)]
    for row in meta_table.rows:
        for i, cell in enumerate(row.cells):
            cell.width = col_widths[i]
            set_cell_background(cell, HEX_BG_LIGHT)
            set_cell_margins(cell, top=80, bottom=80, left=120, right=120)

    start_str = prep_data["start_date"].strftime("%d/%m/%Y") if hasattr(prep_data["start_date"], "strftime") else str(prep_data["start_date"])
    end_str = prep_data["end_date"].strftime("%d/%m/%Y") if hasattr(prep_data["end_date"], "strftime") else str(prep_data["end_date"])

    meta_table.cell(0, 0).paragraphs[0].add_run(f"📅 Período Analizado: {start_str} - {end_str}").bold = True
    meta_table.cell(0, 1).paragraphs[0].add_run(f"🏢 Universo: {len(prep_data['stock_cols'])} Acciones BMV + IPC").bold = True
    meta_table.cell(1, 0).paragraphs[0].add_run(f"🎯 Restricción de Riesgo: σ_p ≤ {opt_data['target_max_std']:.1%} mensual").bold = True
    meta_table.cell(1, 1).paragraphs[0].add_run("🤖 Agente: Cuantitativo Financiero Autónomo").bold = True

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # 3. Resumen Ejecutivo
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Resumen Ejecutivo")
    h1_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_exec = doc.add_paragraph()
    p_exec.paragraph_format.space_after = Pt(10)
    p_exec.paragraph_format.line_spacing = 1.15
    p_exec.add_run(agent_report["executive_summary"])

    # 4. Auditoría y Reglas Metodológicas Críticas
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. Cumplimiento de Reglas Metodológicas Estrictas")
    h2_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_rules = doc.add_paragraph()
    p_rules.paragraph_format.line_spacing = 1.15
    p_rules.add_run("• Regla del Tiempo Crítica: ").bold = True
    p_rules.add_run(agent_report["time_audit"]["text"] + "\n")
    p_rules.add_run("• Regla de No Anualización: ").bold = True
    p_rules.add_run("Todos los cálculos matemáticos (retornos, varianzas, desviaciones estándar y covarianzas) permanecen estrictamente mensuales. No se aplicó el factor 12 ni √12.\n")
    p_rules.add_run("• Exclusión del IPC: ").bold = True
    p_rules.add_run("El índice IPC fue excluido estrictamente de la optimización del portafolio y se utilizó únicamente como benchmark de mercado para Alpha y Beta.\n")
    p_rules.add_run("• Restricciones de Cartera: ").bold = True
    p_rules.add_run(f"100% del capital invertido (Σ w_i = 1), solo posiciones largas (w_i ≥ 0) y riesgo mensual acotado a σ_p ≤ {opt_data['target_max_std']:.1%}.")

    # 5. Tabla de Parámetros Estadísticos Individuales
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. Parámetros Estadísticos Mensuales Individuales")
    h3_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    stats_df = stats_data["stats_table"]
    table_stats = doc.add_table(rows=len(stats_df) + 1, cols=5)
    table_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    table_stats.autofit = False

    headers = ["Activo", "Retorno Mensual", "Volatilidad Mensual", "Varianza Mensual", "Sharpe (Rf=0)"]
    widths = [Inches(1.8), Inches(1.2), Inches(1.2), Inches(1.2), Inches(1.2)]

    # Formatear encabezado
    hdr_row = table_stats.rows[0]
    for i, h_text in enumerate(headers):
        cell = hdr_row.cells[i]
        cell.width = widths[i]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)

    for r_idx, (asset, row) in enumerate(stats_df.iterrows(), start=1):
        row_cells = table_stats.rows[r_idx].cells
        for c_idx in range(5):
            row_cells[c_idx].width = widths[c_idx]
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=80, right=80)
            if r_idx % 2 == 0:
                set_cell_background(row_cells[c_idx], HEX_BG_LIGHT)

        p0 = row_cells[0].paragraphs[0]
        p0.add_run(str(asset)).bold = True
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT

        p1 = row_cells[1].paragraphs[0]
        p1.add_run(f"{row['Rendimiento Mensual Promedio']*100:.2f}%")
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p2 = row_cells[2].paragraphs[0]
        p2.add_run(f"{row['Volatilidad Mensual (Desv. Est.)']*100:.2f}%")
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p3 = row_cells[3].paragraphs[0]
        p3.add_run(f"{row['Varianza Mensual']:.5f}")
        p3.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p4 = row_cells[4].paragraphs[0]
        p4.add_run(f"{row['Ratio Sharpe Mensual (Rf=0)']:.4f}")
        p4.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # 6. Cartera Óptima de Markowitz
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run(f"4. Cartera Óptima de Markowitz (σ_p ≤ {opt_data['target_max_std']*100:.2f}% Mensual)")
    h4_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_opt = doc.add_paragraph()
    p_opt.paragraph_format.line_spacing = 1.15
    p_opt.add_run(f"• Rendimiento Mensual Esperado (μ_p): ").bold = True
    p_opt.add_run(f"{opt_data['expected_return']*100:.3f}%\n")
    p_opt.add_run(f"• Volatilidad Mensual de la Cartera (σ_p): ").bold = True
    p_opt.add_run(f"{opt_data['volatility']*100:.3f}% (Restricción de ≤ {opt_data['target_max_std']*100:.2f}% satisfecha)\n")
    p_opt.add_run(f"• Ratio de Sharpe Mensual (Rf=0): ").bold = True
    p_opt.add_run(f"{opt_data['sharpe_ratio']:.4f}\n")

    # Tabla de Ponderaciones Óptimas
    weights_df = opt_data["weights_df"]
    table_w = doc.add_table(rows=len(weights_df) + 1, cols=3)
    table_w.alignment = WD_TABLE_ALIGNMENT.CENTER
    w_headers = ["Activo", "Ponderación Decimal (w_i)", "Asignación Porcentual (%)"]
    w_widths = [Inches(2.5), Inches(2.0), Inches(2.1)]

    hdr_w = table_w.rows[0]
    for i, h_text in enumerate(w_headers):
        cell = hdr_w.cells[i]
        cell.width = w_widths[i]
        set_cell_background(cell, HEX_NAVY)
        set_cell_margins(cell, top=100, bottom=100, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(h_text)
        run.bold = True
        run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        run.font.size = Pt(9)

    for r_idx, (_, row) in enumerate(weights_df.iterrows(), start=1):
        row_cells = table_w.rows[r_idx].cells
        for c_idx in range(3):
            row_cells[c_idx].width = w_widths[c_idx]
            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=80, right=80)
            if r_idx % 2 == 0:
                set_cell_background(row_cells[c_idx], HEX_BG_LIGHT)

        p0 = row_cells[0].paragraphs[0]
        p0.add_run(str(row["Activo"])).bold = True
        p0.alignment = WD_ALIGN_PARAGRAPH.LEFT

        p1 = row_cells[1].paragraphs[0]
        p1.add_run(f"{row['Ponderación Óptima (w)']:.4f}")
        p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT

        p2 = row_cells[2].paragraphs[0]
        w_pct = row['Porcentaje (%)']
        r_pct = p2.add_run(f"{w_pct:.2f}%")
        if w_pct > 0.01:
            r_pct.bold = True
        p2.alignment = WD_ALIGN_PARAGRAPH.RIGHT

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # 7. Generar e incrustar gráficos
    charts = create_charts_images(stats_data, opt_data, mc_data, bench_data)

    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run("5. Simulación de Monte Carlo y Asignación Gráfica")
    h5_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_chart1 = doc.add_paragraph()
    p_chart1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(charts["monte_carlo"], width=Inches(6.2))
    cap1 = doc.add_paragraph("Figura 1: Nube de 1,000 carteras aleatorias de Monte Carlo, Cartera Óptima y Benchmark.")
    cap1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap1.style.font.size = Pt(8.5)
    cap1.style.font.italic = True

    p_chart2 = doc.add_paragraph()
    p_chart2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(charts["weights"], width=Inches(6.0))
    cap2 = doc.add_paragraph("Figura 2: Distribución de pesos óptimos del portafolio en posiciones largas.")
    cap2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap2.style.font.size = Pt(8.5)
    cap2.style.font.italic = True

    p_chart3 = doc.add_paragraph()
    p_chart3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_picture(charts["correlation"], width=Inches(5.8))
    cap3 = doc.add_paragraph("Figura 3: Matriz de calor de correlaciones mensuales entre activos.")
    cap3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap3.style.font.size = Pt(8.5)
    cap3.style.font.italic = True

    # 8. Comparativa contra el Benchmark (IPC)
    if bench_data:
        h6 = doc.add_heading(level=1)
        h6_run = h6.add_run("6. Evaluación frente al Benchmark de Mercado (IPC)")
        h6_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

        p_bench = doc.add_paragraph()
        p_bench.paragraph_format.line_spacing = 1.15
        p_bench.add_run(agent_report["bench_insights"]["text"])

    # 9. Diagnóstico Exhaustivo del Agente Cuantitativo Autónomo
    h7 = doc.add_heading(level=1)
    h7_run = h7.add_run("7. Dictamen del Agente Cuantitativo Financiero Autónomo")
    h7_run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)

    p_diag = doc.add_paragraph()
    p_diag.paragraph_format.line_spacing = 1.15
    p_diag.add_run(agent_report["opt_insights"]["text"] + "\n\n")
    p_diag.add_run(agent_report["mc_insights"]["text"] + "\n\n")
    p_diag.add_run(agent_report["recommendations"])

    # Guardar documento
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)

    if output_path:
        with open(output_path, "wb") as f:
            f.write(buf.getvalue())

    return buf
