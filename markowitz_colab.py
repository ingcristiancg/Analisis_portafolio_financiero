"""
================================================================================
SCRIPT CUANTITATIVO: OPTIMIZACIÓN DE CARTERAS DE MARKOWITZ Y MONTE CARLO
Diseñado para ejecución en Google Colab o Entorno Local
================================================================================
Reglas y Pasos de Cálculo Estrictos:
1️⃣ Preprocesamiento de Datos (REGLA DEL TIEMPO CRÍTICA):
   - Orden cronológico garantizado (pasado -> presente) antes de rendimientos.
   - Si el archivo viene con fechas descendentes, se invierte.
   - Rendimientos simples mensuales: R_t = (P_t - P_{t-1}) / P_{t-1}.
   - Exclusión del IPC en la optimización.
   - Eliminación de nulos (NaN).
2️⃣ Parámetros Estadísticos (ESTRICTAMENTE MENSUALES):
   - Medias aritméticas mensuales, matriz de covarianza y correlación.
   - PROHIBIDO ANUALIZAR (no multiplicar por 12 ni sqrt(12)).
3️⃣ Optimización de la Cartera (Frontera Eficiente):
   - Maximizar rendimiento esperado mensual de la cartera.
   - Restricción 1: Desviación estándar mensual sigma_p <= 0.07.
   - Restricción 2: 100% del capital invertido (sum(w_i) = 1).
   - Límites: Solo largos (w_i >= 0).
4️⃣ Simulación de Monte Carlo:
   - 1,000 carteras aleatorias respetando sum(w_i)=1 y w_i >= 0.
   - Rendimiento, riesgo y Ratio Sharpe (Rf=0) mensuales.
   - Gráfico de la nube de la frontera eficiente.
5️⃣ Generación del Dashboard & Exportación a Word (.docx).
================================================================================
"""

import os
import sys
import io
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import seaborn as sns
from scipy.optimize import minimize

# Intentar importar python-docx para exportar el reporte
try:
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls
except ImportError:
    print("Instalando python-docx...")
    os.system("pip install python-docx")
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import parse_xml
    from docx.oxml.ns import nsdecls

# Detectar si estamos en Google Colab
IN_COLAB = "google.colab" in sys.modules

# ==============================================================================
# PASO 1: PREPROCESAMIENTO DE DATOS (REGLA DEL TIEMPO CRÍTICA)
# ==============================================================================
def cargar_y_preprocesar_excel(ruta_o_archivo, hoja=None, col_ipc="IPC"):
    print("\n" + "="*70)
    print("1️⃣ PREPROCESAMIENTO DE DATOS (REGLA DEL TIEMPO CRÍTICA)")
    print("="*70)

    # 1. Leer Excel
    xl = pd.ExcelFile(ruta_o_archivo)
    if hoja is None:
        hoja = "report" if "report" in xl.sheet_names else xl.sheet_names[0]
    df_raw = pd.read_excel(ruta_o_archivo, sheet_name=hoja)
    print(f"✓ Archivo cargado exitosamente. Hoja seleccionada: '{hoja}'")

    # 2. Identificar columna de fechas
    col_fecha = None
    for c in df_raw.columns:
        if any(k in str(c).lower() for k in ["fecha", "date", "periodo", "mes"]):
            col_fecha = c
            break
    if col_fecha is None:
        col_fecha = df_raw.columns[0]

    # 3. Limpiar filas no numéricas o pies de página (Promedio, Varianza, etc.)
    df_clean = df_raw.copy()
    df_clean["__fecha__"] = pd.to_datetime(df_clean[col_fecha], errors="coerce")
    df_clean = df_clean.dropna(subset=["__fecha__"]).copy()

    # 4. REGLA DEL TIEMPO CRÍTICA:
    # Si viene ordenado de la fecha más reciente (arriba) a la más antigua (abajo),
    # DEBES invertir el orden para que el dato más antiguo quede en la primera fila.
    primera_fecha = df_clean["__fecha__"].iloc[0]
    ultima_fecha = df_clean["__fecha__"].iloc[-1]
    
    if primera_fecha > ultima_fecha:
        print("⚠️ ALERTA DE CRONOLOGÍA: El archivo viene ordenado de más reciente a más antiguo.")
        print("⏳ INVIRTIENDO EL ORDEN: El dato más antiguo se posiciona en la fila 0 (pasado -> presente).")
        df_clean = df_clean.iloc[::-1].reset_index(drop=True)
    else:
        print("✅ CRONOLOGÍA VALIDADA: El archivo ya está ordenado de pasado a presente.")
        df_clean = df_clean.sort_values(by="__fecha__", ascending=True).reset_index(drop=True)

    # 5. Filtrar columnas numéricas válidas (excluir columnas de fórmulas de Excel como Unnamed o Proporción)
    columnas_validas = []
    for c in df_clean.columns:
        if c in [col_fecha, "__fecha__"] or str(c).startswith("Unnamed:") or "proporci" in str(c).lower() or "portafolio" in str(c).lower():
            continue
        s = pd.to_numeric(df_clean[c], errors="coerce")
        if s.notna().sum() > len(df_clean) * 0.5:
            df_clean[c] = s
            columnas_validas.append(c)

    # Separar IPC de las acciones a optimizar
    col_ipc_real = None
    for c in columnas_validas:
        if col_ipc.lower() in str(c).lower():
            col_ipc_real = c
            break

    acciones = [c for c in columnas_validas if c != col_ipc_real]
    print(f"✓ Activos detectados para optimización ({len(acciones)}): {acciones}")
    if col_ipc_real:
        print(f"✓ Columna del Benchmark detectada: '{col_ipc_real}' (EXCLUIDA de la optimización de la cartera).")

    # 6. Rendimientos simples mensuales: R_t = (P_t - P_{t-1}) / P_{t-1}
    rendimientos = df_clean[columnas_validas].pct_change().dropna(how="all")
    rendimientos.index = df_clean["__fecha__"].iloc[1:]
    rendimientos = rendimientos.dropna(subset=acciones)

    rend_acciones = rendimientos[acciones]
    rend_ipc = rendimientos[col_ipc_real] if col_ipc_real else None

    print(f"✓ Períodos mensuales calculados: {len(rend_acciones)} meses (desde {rendimientos.index[0].strftime('%b %Y')} hasta {rendimientos.index[-1].strftime('%b %Y')}).")

    return {
        "df_precios": df_clean,
        "rendimientos_acciones": rend_acciones,
        "rendimientos_ipc": rend_ipc,
        "acciones": acciones,
        "col_ipc": col_ipc_real,
        "fechas": pd.Series(rend_acciones.index, index=rend_acciones.index),
    }


# ==============================================================================
# PASO 2: PARÁMETROS ESTADÍSTICOS (ESTRICTAMENTE MENSUALES)
# ==============================================================================
def calcular_estadisticas_mensuales(rend_acciones):
    print("\n" + "="*70)
    print("2️⃣ PARÁMETROS ESTADÍSTICOS (ESTRICTAMENTE MENSUALES)")
    print("🚨 REGLA CRÍTICA ABSOLUTA: PROHIBIDO ANUALIZAR. TODO EL CÁLCULO ES MENSUAL.")
    print("="*70)

    # Media aritmética mensual de cada acción
    media_mensual = rend_acciones.mean()
    # Desviación estándar mensual
    std_mensual = rend_acciones.std()
    # Varianza mensual
    var_mensual = rend_acciones.var()
    # Matriz de covarianza mensual
    matriz_cov = rend_acciones.cov()
    # Matriz de correlación
    matriz_corr = rend_acciones.corr()
    # Ratio de Sharpe mensual individual (Rf = 0)
    sharpe_mensual = media_mensual / std_mensual

    tabla_stats = pd.DataFrame({
        "Rendimiento Mensual": media_mensual,
        "Volatilidad Mensual": std_mensual,
        "Varianza Mensual": var_mensual,
        "Ratio Sharpe (Rf=0)": sharpe_mensual,
    })

    print("\nTabla de Parámetros Estadísticos Mensuales:")
    print(tabla_stats.to_string(formatters={
        "Rendimiento Mensual": "{:.2%}".format,
        "Volatilidad Mensual": "{:.2%}".format,
        "Varianza Mensual": "{:.6f}".format,
        "Ratio Sharpe (Rf=0)": "{:.4f}".format,
    }))

    return {
        "media_mensual": media_mensual,
        "std_mensual": std_mensual,
        "matriz_cov": matriz_cov,
        "matriz_corr": matriz_corr,
        "tabla_stats": tabla_stats,
    }


# ==============================================================================
# PASO 3: OPTIMIZACIÓN DE LA CARTERA (FRONTERA EFICIENTE DE MARKOWITZ)
# ==============================================================================
def optimizar_cartera_markowitz(rend_acciones, max_std=0.07):
    print("\n" + "="*70)
    print(f"3️⃣ OPTIMIZACIÓN DE LA CARTERA (Restricción σ_p ≤ {max_std:.2%})")
    print("="*70)

    mu = rend_acciones.mean().values
    cov = rend_acciones.cov().values
    n = len(mu)
    nombres = rend_acciones.columns.tolist()

    # Función objetivo: Maximizar rendimiento mensual -> Minimizar -E[R]
    def obj_neg_return(w):
        return -float(np.dot(w, mu))

    def portfolio_std(w):
        return float(np.sqrt(np.dot(w.T, np.dot(cov, w))))

    # Restricciones:
    # 1. 100% del capital invertido: sum(w_i) = 1
    # 2. Desviación estándar mensual <= 0.07: 0.07 - sigma_p >= 0
    restricciones = [
        {"type": "eq", "fun": lambda w: np.sum(w) - 1.0},
        {"type": "ineq", "fun": lambda w: max_std - np.sqrt(np.dot(w.T, np.dot(cov, w)))},
    ]

    # Límites: solo posiciones largas (w_i >= 0)
    limites = tuple((0.0, 1.0) for _ in range(n))
    w0 = np.ones(n) / n

    res = minimize(obj_neg_return, w0, method="SLSQP", bounds=limites, constraints=restricciones)

    if not res.success:
        raise RuntimeError(f"Error en la optimización: {res.message}")

    pesos_optimos = res.x
    pesos_optimos = np.where(pesos_optimos < 1e-5, 0.0, pesos_optimos)
    pesos_optimos = pesos_optimos / np.sum(pesos_optimos)

    retorno_optimo = float(np.dot(pesos_optimos, mu))
    volatilidad_optima = portfolio_std(pesos_optimos)
    sharpe_optimo = retorno_optimo / volatilidad_optima if volatilidad_optima > 0 else 0.0

    df_pesos = pd.DataFrame({
        "Activo": nombres,
        "Ponderación Decimal (w)": pesos_optimos,
        "Porcentaje (%)": pesos_optimos * 100.0,
    }).sort_values(by="Ponderación Decimal (w)", ascending=False).reset_index(drop=True)

    print(f"🎯 Rendimiento Esperado Mensual: {retorno_optimo*100:.3f}%")
    print(f"🛡️ Volatilidad Mensual (σ_p): {volatilidad_optima*100:.3f}% (Restricción ≤ {max_std*100:.1f}% CUMPLIDA)")
    print(f"⚖️ Ratio de Sharpe Mensual (Rf=0): {sharpe_optimo:.4f}")
    print("\nAsignación Óptima de Capital (Posiciones Largas):")
    for _, fila in df_pesos[df_pesos["Ponderación Decimal (w)"] > 0.001].iterrows():
        print(f"  • {fila['Activo']}: {fila['Porcentaje (%)']:.2f}%")

    retornos_historicos_cartera = rend_acciones.dot(pesos_optimos)

    return {
        "pesos": pesos_optimos,
        "df_pesos": df_pesos,
        "retorno_optimo": retorno_optimo,
        "volatilidad_optima": volatilidad_optima,
        "sharpe_optimo": sharpe_optimo,
        "retornos_historicos": retornos_historicos_cartera,
    }


# ==============================================================================
# PASO 4: SIMULACIÓN DE MONTE CARLO (1,000 CARTERAS)
# ==============================================================================
def simular_monte_carlo(rend_acciones, num_carteras=1000, semilla=42):
    print("\n" + "="*70)
    print(f"4️⃣ SIMULACIÓN DE MONTE CARLO ({num_carteras} Carteras Aleatorias)")
    print("="*70)

    np.random.seed(semilla)
    mu = rend_acciones.mean().values
    cov = rend_acciones.cov().values
    n = len(mu)

    retornos_sim = np.zeros(num_carteras)
    volatilidades_sim = np.zeros(num_carteras)
    sharpes_sim = np.zeros(num_carteras)

    for i in range(num_carteras):
        w = np.random.exponential(scale=1.0, size=n)
        w = w / np.sum(w)
        
        r = np.dot(w, mu)
        vol = np.sqrt(np.dot(w.T, np.dot(cov, w)))
        s = r / vol if vol > 0 else 0.0

        retornos_sim[i] = r
        volatilidades_sim[i] = vol
        sharpes_sim[i] = s

    df_mc = pd.DataFrame({
        "Rendimiento Mensual": retornos_sim,
        "Volatilidad Mensual": volatilidades_sim,
        "Ratio Sharpe": sharpes_sim,
    })

    print(f"✓ Simulación completada para {num_carteras} carteras.")
    print(f"  • Retorno medio simulado: {retornos_sim.mean()*100:.2f}%")
    print(f"  • Volatilidad media simulada: {volatilidades_sim.mean()*100:.2f}%")
    print(f"  • Máximo Sharpe simulado: {sharpes_sim.max():.4f}")

    return df_mc


# ==============================================================================
# PASO 5: DASHBOARD GRÁFICO & EXPORTACIÓN A WORD
# ==============================================================================
def generar_graficos_y_dashboard(df_mc, opt_res, stats_res, rend_ipc=None):
    print("\n" + "="*70)
    print("5️⃣ GENERACIÓN DE GRÁFICOS DEL DASHBOARD")
    print("="*70)

    sns.set_theme(style="whitegrid")
    fig, axes = plt.subplots(1, 2, figsize=(15, 6), dpi=150)

    # Gráfico 1: Nube de Monte Carlo
    ax1 = axes[0]
    sc = ax1.scatter(
        df_mc["Volatilidad Mensual"] * 100,
        df_mc["Rendimiento Mensual"] * 100,
        c=df_mc["Ratio Sharpe"],
        cmap="viridis",
        alpha=0.6,
        s=20,
    )
    cbar = plt.colorbar(sc, ax=ax1)
    cbar.set_label("Ratio de Sharpe Mensual (Rf=0)")

    opt_ret = opt_res["retorno_optimo"] * 100
    opt_vol = opt_res["volatilidad_optima"] * 100

    ax1.scatter([opt_vol], [opt_ret], color="red", s=180, marker="*", label="Cartera Óptima Markowitz", zorder=5)
    ax1.axvline(x=7.0, color="darkred", linestyle="--", label="Restricción Riesgo (≤ 7.0%)")

    if rend_ipc is not None:
        ipc_ret = rend_ipc.mean() * 100
        ipc_vol = rend_ipc.std() * 100
        ax1.scatter([ipc_vol], [ipc_ret], color="green", s=120, marker="D", label="Benchmark IPC", zorder=4)

    ax1.set_title("Nube de Monte Carlo & Cartera Óptima (Datos Mensuales)")
    ax1.set_xlabel("Volatilidad Mensual σ_p (%)")
    ax1.set_ylabel("Rendimiento Mensual μ_p (%)")
    ax1.legend(loc="upper left")

    # Gráfico 2: Asignación de Pesos
    ax2 = axes[1]
    df_activos = opt_res["df_pesos"][opt_res["df_pesos"]["Porcentaje (%)"] > 0.1]
    ax2.barh(df_activos["Activo"], df_activos["Porcentaje (%)"], color="#1D4ED8", height=0.55)
    ax2.set_title("Ponderación Óptima de Activos (Markowitz Long-Only)")
    ax2.set_xlabel("Ponderación en Cartera (%)")
    ax2.invert_yaxis()
    for i, v in enumerate(df_activos["Porcentaje (%)"]):
        ax2.text(v + 0.5, i, f"{v:.1f}%", va="center", fontweight="bold")

    plt.tight_layout()
    plt.savefig("dashboard_markowitz.png", bbox_inches="tight")
    if IN_COLAB:
        plt.show()
    else:
        try:
            plt.show(block=False)
            plt.pause(0.5)
        except Exception:
            pass

    return fig


def exportar_informe_word(prep_res, stats_res, opt_res, df_mc, ruta_salida="Informe_Markowitz_Colab.docx"):
    print("\n" + "="*70)
    print("📄 EXPORTACIÓN DEL INFORME EJECUTIVO A WORD (.DOCX)")
    print("="*70)

    doc = Document()
    
    # Título
    p_title = doc.add_heading(level=0)
    p_title.add_run("INFORME EJECUTIVO: OPTIMIZACIÓN DE CARTERAS DE MARKOWITZ")

    p_sub = doc.add_paragraph("Análisis Cuantitativo Estricto Mensual y Simulación de Monte Carlo")
    p_sub.style.font.italic = True

    # Resumen
    doc.add_heading("1. Resumen de la Cartera Óptima", level=1)
    doc.add_paragraph(
        f"• Rendimiento Mensual Esperado: {opt_res['retorno_optimo']*100:.3f}%\n"
        f"• Volatilidad Mensual (σ_p): {opt_res['volatilidad_optima']*100:.3f}% (Restricción ≤ 7.00% cumplida)\n"
        f"• Ratio de Sharpe Mensual (Rf=0): {opt_res['sharpe_optimo']:.4f}\n"
        f"• Total períodos analizados: {len(prep_res['rendimientos_acciones'])} meses."
    )

    # Tabla de Pesos
    doc.add_heading("2. Asignación Óptima de Capital", level=1)
    t = doc.add_table(rows=1, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    hdr[0].text = "Activo"
    hdr[1].text = "Ponderación (w)"
    hdr[2].text = "Porcentaje (%)"

    for _, row in opt_res["df_pesos"].iterrows():
        r = t.add_row().cells
        r[0].text = str(row["Activo"])
        r[1].text = f"{row['Ponderación Decimal (w)']:.4f}"
        r[2].text = f"{row['Porcentaje (%)']:.2f}%"

    # Parámetros estadísticos
    doc.add_heading("3. Parámetros Estadísticos Mensuales Individuales", level=1)
    t_stats = doc.add_table(rows=1, cols=4)
    t_stats.alignment = WD_TABLE_ALIGNMENT.CENTER
    h_s = t_stats.rows[0].cells
    h_s[0].text = "Activo"
    h_s[1].text = "Rendimiento Mensual"
    h_s[2].text = "Volatilidad Mensual"
    h_s[3].text = "Sharpe (Rf=0)"

    for activo, row in stats_res["tabla_stats"].iterrows():
        r = t_stats.add_row().cells
        r[0].text = str(activo)
        r[1].text = f"{row['Rendimiento Mensual']*100:.2f}%"
        r[2].text = f"{row['Volatilidad Mensual']*100:.2f}%"
        r[3].text = f"{row['Ratio Sharpe (Rf=0)']:.4f}"

    doc.save(ruta_salida)
    print(f"✅ Documento Word guardado exitosamente en: '{ruta_salida}'")

    if IN_COLAB:
        try:
            from google.colab import files
            files.download(ruta_salida)
            print("✓ Descarga automática en Google Colab iniciada.")
        except Exception as e:
            print(f"Descarga manual disponible: {ruta_salida}")


# ==============================================================================
# EJECUCIÓN PRINCIPAL
# ==============================================================================
if __name__ == "__main__":
    archivo_excel = "Histórico_Acciones.xlsx"
    
    if not os.path.exists(archivo_excel):
        if IN_COLAB:
            print("Por favor sube el archivo de Excel histórico:")
            from google.colab import files
            uploaded = files.upload()
            archivo_excel = list(uploaded.keys())[0]
        else:
            print(f"Error: No se encontró el archivo '{archivo_excel}'.")
            sys.exit(1)

    # 1. Preprocesamiento
    prep = cargar_y_preprocesar_excel(archivo_excel)

    # 2. Estadísticas mensuales
    stats = calcular_estadisticas_mensuales(prep["rendimientos_acciones"])

    # 3. Optimización Markowitz (sigma <= 0.07)
    opt = optimizar_cartera_markowitz(prep["rendimientos_acciones"], max_std=0.07)

    # 4. Monte Carlo (1,000 carteras)
    mc = simular_monte_carlo(prep["rendimientos_acciones"], num_carteras=1000)

    # 5. Gráficos y Word
    generar_graficos_y_dashboard(mc, opt, stats, prep["rendimientos_ipc"])
    exportar_informe_word(prep, stats, opt, mc, ruta_salida="Informe_Markowitz_Colab.docx")
    print("\n🎉 Proceso completado con éxito cumpliendo todas las reglas estrictas.")
